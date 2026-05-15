"""
Generates the project presentation PPT — clean, general, not overly technical.
Run: python generate_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Colour palette ────────────────────────────────────────────────────────────
NAVY      = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT    = RGBColor(0x16, 0x21, 0x3E)
HIGHLIGHT = RGBColor(0xE9, 0x4F, 0x37)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT     = RGBColor(0xF0, 0xF0, 0xF0)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x2E)
MID_GREY  = RGBColor(0x99, 0x99, 0x99)
CARD_BG   = RGBColor(0x25, 0x25, 0x4A)
LIGHT_BG  = RGBColor(0xF5, 0xF5, 0xF5)
CARD_LITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ── Helpers ───────────────────────────────────────────────────────────────────

def rect(slide, l, t, w, h, fill=None, line=None, lw=Pt(0)):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.line.width = lw
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line
    else:
        s.line.fill.background()
    return s


def text(slide, content, l, t, w, h,
         size=16, bold=False, color=WHITE,
         align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = content
    r.font.size   = Pt(size)
    r.font.bold   = bold
    r.font.color.rgb = color
    r.font.italic = italic
    return txb


def dark_slide(title):
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, 13.33, 7.5, fill=NAVY)
    rect(slide, 0, 0, 0.07, 7.5, fill=HIGHLIGHT)
    text(slide, title, 0.35, 0.22, 12.5, 0.8, size=30, bold=True, color=WHITE)
    rect(slide, 0.35, 1.05, 1.4, 0.05, fill=HIGHLIGHT)
    return slide


def light_slide(title):
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, 13.33, 7.5, fill=LIGHT_BG)
    rect(slide, 0, 0, 13.33, 1.15, fill=NAVY)
    rect(slide, 0, 1.15, 13.33, 0.05, fill=HIGHLIGHT)
    text(slide, title, 0.4, 0.2, 12.5, 0.75, size=30, bold=True, color=WHITE)
    return slide


def bullet_block(slide, items, l, t, w, h,
                 size=15, color=LIGHT, gap=Pt(18), dark=True):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = True
    tf = txb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = gap
        r = p.add_run()
        r.text = item
        r.font.size = Pt(size)
        r.font.color.rgb = color


def simple_card(slide, l, t, w, h, heading, lines, dark=True):
    bg   = CARD_BG   if dark else CARD_LITE
    tc   = HIGHLIGHT
    body = LIGHT     if dark else DARK_TEXT
    rect(slide, l, t, w, h, fill=bg, line=HIGHLIGHT, lw=Pt(1))
    text(slide, heading, l+0.18, t+0.15, w-0.36, 0.5,
         size=14, bold=True, color=tc)
    rect(slide, l+0.18, t+0.65, w-0.36, 0.03, fill=HIGHLIGHT)
    bullet_block(slide, lines, l+0.18, t+0.75, w-0.36, h-0.9,
                 size=12, color=body, gap=Pt(10), dark=dark)


# =============================================================================
# SLIDE 1 — TITLE
# =============================================================================
slide = prs.slides.add_slide(BLANK)
rect(slide, 0, 0, 13.33, 7.5, fill=ACCENT)
rect(slide, 0, 0, 0.08, 7.5, fill=HIGHLIGHT)
rect(slide, 0, 5.6, 13.33, 1.9, fill=NAVY)

text(slide, "E-Commerce AI Support Platform",
     0.55, 1.5, 12.0, 1.3, size=38, bold=True, color=WHITE)
rect(slide, 0.55, 2.85, 3.0, 0.06, fill=HIGHLIGHT)
text(slide, "Automated customer support powered by LLM, RAG, and real e-commerce data",
     0.55, 2.98, 11.5, 0.65, size=17, color=LIGHT)

text(slide, "Groq  |  ChromaDB  |  Streamlit  |  SQLite  |  Sentence Transformers",
     0.55, 3.72, 11.0, 0.5, size=13, color=MID_GREY, italic=True)

text(slide, "Muhammad Huzaifa  |  2025",
     0.55, 6.2, 8.0, 0.45, size=12, color=MID_GREY)


# =============================================================================
# SLIDE 2 — PROBLEM STATEMENT
# =============================================================================
slide = dark_slide("Problem Statement")

problems = [
    ("High Volume",      "E-commerce businesses receive thousands of customer queries every day about orders, products, and refunds."),
    ("High Cost",        "Running a large human support team is expensive and cannot scale with traffic spikes or off-hours demand."),
    ("Inconsistency",    "Different agents give different answers. Response quality varies and there is no single source of truth."),
    ("Missed Insights",  "Support interactions contain valuable data about recurring issues, but it is rarely captured or analysed."),
]

y = 1.25
for label, body in problems:
    rect(slide, 0.35, y, 12.6, 1.25, fill=CARD_BG)
    text(slide, label, 0.55, y+0.12, 2.5, 0.5,
         size=14, bold=True, color=HIGHLIGHT)
    text(slide, body, 3.1, y+0.16, 9.7, 0.85,
         size=13, color=LIGHT, wrap=True)
    y += 1.42


# =============================================================================
# SLIDE 3 — SOLUTION
# =============================================================================
slide = dark_slide("Our Solution")

text(slide, "An AI-powered customer support assistant that handles queries automatically,\ndraws on real product and order data, and generates natural responses using a large language model.",
     0.35, 1.2, 12.6, 1.1, size=15, color=LIGHT, wrap=True)

pillars = [
    ("Conversational AI",  "Understands customer questions, asks follow-up questions when needed, and maintains context throughout the conversation."),
    ("Real Data Queries",  "Fetches live order status, product prices, inventory, and discounts directly from the database — no hardcoded answers."),
    ("LLM Responses",      "Uses the Groq large language model to generate natural, helpful replies grounded in retrieved policy documents."),
    ("User Accounts",      "Customers can register, log in, and access their full conversation and support ticket history at any time."),
]

col_positions = [0.35, 3.6, 6.85, 10.1]
for i, (heading, body) in enumerate(pillars):
    l = col_positions[i]
    rect(slide, l, 2.55, 3.0, 4.55, fill=CARD_BG, line=HIGHLIGHT, lw=Pt(1))
    text(slide, heading, l+0.18, 2.68, 2.65, 0.55, size=13, bold=True, color=HIGHLIGHT)
    rect(slide, l+0.18, 3.23, 2.65, 0.03, fill=HIGHLIGHT)
    text(slide, body, l+0.18, 3.32, 2.65, 3.6, size=12, color=LIGHT, wrap=True)


# =============================================================================
# SLIDE 4 — DATASETS
# =============================================================================
slide = light_slide("Datasets")

text(slide, "Two dataset sources — each serving a distinct role in the system.",
     0.4, 1.28, 12.5, 0.45, size=14, color=DARK_TEXT, italic=True)

# Left — Pakistani / Daraz
rect(slide, 0.35, 1.85, 7.9, 5.3, fill=NAVY, line=HIGHLIGHT, lw=Pt(1))
text(slide, "AI Assistant Database", 0.55, 1.98, 7.5, 0.5,
     size=13, bold=True, color=HIGHLIGHT)
rect(slide, 0.55, 2.48, 7.5, 0.03, fill=HIGHLIGHT)

daraz_rows = [
    ("Pakistan E-Commerce Dataset",   "Orders, order history, payment methods"),
    ("Daraz Top Sales Products 2024", "Product catalog with PKR prices and ratings"),
    ("Daraz 11.11 Top Selling Data",  "Discounted and sale products"),
]
y = 2.6
for name, role in daraz_rows:
    rect(slide, 0.45, y, 7.7, 1.4, fill=CARD_BG)
    text(slide, name, 0.65, y+0.1,  4.5, 0.5, size=13, bold=True, color=WHITE)
    text(slide, role, 0.65, y+0.62, 7.1, 0.55, size=12, color=LIGHT)
    y += 1.5

# Right — Olist
rect(slide, 8.65, 1.85, 4.35, 5.3, fill=NAVY, line=MID_GREY, lw=Pt(1))
text(slide, "Analytics Dashboard", 8.85, 1.98, 3.95, 0.5,
     size=13, bold=True, color=MID_GREY)
rect(slide, 8.85, 2.48, 3.95, 0.03, fill=MID_GREY)
text(slide, "Olist Brazilian\nE-Commerce Dataset", 8.85, 2.6, 3.95, 0.85,
     size=15, bold=True, color=WHITE)
text(slide, "100,000+ real orders\nfrom 2016 to 2018", 8.85, 3.5, 3.95, 0.65,
     size=13, color=LIGHT)
text(slide, "Used for generating dashboard charts, delivery delay analysis, review score trends, and operational KPIs. Not used by the AI assistant.",
     8.85, 4.3, 3.95, 2.5, size=12, color=LIGHT, wrap=True)


# =============================================================================
# SLIDE 5 — HOW IT WORKS
# =============================================================================
slide = dark_slide("How It Works")

# Flow boxes
steps = [
    ("1. Customer sends a message",
     "The system reads the message and identifies what the customer is asking about."),
    ("2. Intent is classified",
     "A sentence embedding model compares the message against known query types to determine the intent."),
    ("3. Information is collected",
     "If the query needs more details (e.g. an order ID), the bot asks follow-up questions one at a time."),
    ("4. Data is retrieved",
     "The system queries the database for order or product facts, and searches policy documents for relevant guidance."),
    ("5. Response is generated",
     "The retrieved information is sent to the Groq LLM, which produces a clear, natural language reply."),
]

y = 1.2
for heading, body in steps:
    rect(slide, 0.35, y, 12.6, 1.08, fill=CARD_BG)
    text(slide, heading, 0.55, y+0.1, 4.0, 0.45,
         size=13, bold=True, color=HIGHLIGHT)
    text(slide, body, 4.55, y+0.14, 8.2, 0.75,
         size=13, color=LIGHT, wrap=True)
    y += 1.22


# =============================================================================
# SLIDE 6 — KEY FEATURES
# =============================================================================
slide = light_slide("Key Features")

features = [
    ("Order Tracking",        ["Check live order status", "Supports all order ID formats", "Shows estimated delivery info"]),
    ("Product Search",        ["Search by name or category", "Filter by budget", "Handles partial / misspelled names"]),
    ("Recommendations",       ["Budget-based suggestions", "Category filtering", "Paginated results on request"]),
    ("Refunds & Cancellation",["Checks order eligibility", "Updates the database on confirmation", "Explains policy inline"]),
    ("Support Tickets",       ["Raises tickets through chat", "Six issue categories", "Linked to user account"]),
    ("User Accounts",         ["Register and log in", "Chat and ticket history saved", "Session persists after refresh"]),
]

col_positions = [0.35, 4.52, 8.69]
row_positions = [1.3, 4.2]

for i, (heading, points) in enumerate(features):
    col = i % 3
    row = i // 3
    l = col_positions[col]
    t = row_positions[row]
    simple_card(slide, l, t, 3.8, 2.65, heading,
                ["- " + p for p in points], dark=False)


# =============================================================================
# SLIDE 7 — TECHNOLOGY STACK
# =============================================================================
slide = dark_slide("Technology Stack")

stack = [
    ("Streamlit",             "Web application framework",     "User interface and dashboard"),
    ("Groq / Llama 3.3-70B",  "Large language model API",     "Natural language response generation"),
    ("ChromaDB",              "Vector database",               "Semantic search over policy documents"),
    ("Sentence Transformers", "Embedding model",               "Intent classification"),
    ("SQLite",                "Relational database",           "Orders, products, users, tickets"),
    ("Pandas",                "Data processing library",       "Dataset loading and dashboard analytics"),
    ("RapidFuzz",             "Fuzzy string matching",         "Product name search"),
    ("TextBlob",              "NLP library",                   "Customer sentiment analysis"),
]

headers = ["Technology", "Type", "Used For"]
col_w = [3.5, 3.5, 5.6]
col_l = [0.35, 3.9, 7.45]

for j, (h, w, l) in enumerate(zip(headers, col_w, col_l)):
    rect(slide, l, 1.2, w, 0.45, fill=HIGHLIGHT)
    text(slide, h, l+0.12, 1.24, w-0.24, 0.38,
         size=13, bold=True, color=WHITE)

for i, (tech, kind, role) in enumerate(stack):
    bg = CARD_BG if i % 2 == 0 else NAVY
    y = 1.65 + i * 0.6
    for val, w, l in zip([tech, kind, role], col_w, col_l):
        rect(slide, l, y, w, 0.58, fill=bg)
        text(slide, val, l+0.12, y+0.1, w-0.24, 0.42,
             size=12, color=WHITE if val == tech else LIGHT, bold=(val == tech))


# =============================================================================
# SLIDE 8 — DASHBOARD OVERVIEW
# =============================================================================
slide = light_slide("Dashboard Overview")

tabs = [
    ("AI Assistant",
     "The main chat interface. Customers type questions and receive instant responses. "
     "Each message shows the detected intent, sentiment, and which data source was used."),
    ("Executive Dashboard",
     "Visual analytics for platform managers. Shows order volume trends, delivery performance, "
     "and top product categories using charts and graphs."),
    ("Support Insights",
     "Operational KPIs including average review scores, delay rates, problematic product categories, "
     "recent chat logs, and a full support ticket breakdown."),
    ("My Account",
     "Personal view for logged-in users. Shows their chat history, raised support tickets, "
     "and account details — all stored and retrieved from the database."),
]

y = 1.3
for i, (tab, desc) in enumerate(tabs):
    rect(slide, 0.35, y, 12.6, 1.35, fill=NAVY)
    rect(slide, 0.35, y, 0.06, 1.35, fill=HIGHLIGHT)
    text(slide, f"Tab {i+1}   {tab}", 0.55, y+0.12, 4.2, 0.5,
         size=14, bold=True, color=WHITE)
    text(slide, desc, 4.75, y+0.22, 8.0, 0.9,
         size=13, color=LIGHT, wrap=True)
    y += 1.5


# =============================================================================
# SLIDE 9 — CONCLUSION
# =============================================================================
slide = prs.slides.add_slide(BLANK)
rect(slide, 0, 0, 13.33, 7.5, fill=ACCENT)
rect(slide, 0, 0, 0.08, 7.5, fill=HIGHLIGHT)
rect(slide, 0, 5.5, 13.33, 2.0, fill=NAVY)

text(slide, "Conclusion", 0.55, 0.75, 12.0, 0.9,
     size=36, bold=True, color=WHITE)
rect(slide, 0.55, 1.65, 2.2, 0.06, fill=HIGHLIGHT)

summary = [
    "A complete customer support system built on real e-commerce data — not mock or demo responses.",
    "Combines database queries, semantic search, and a large language model for accurate, natural answers.",
    "Handles the full support lifecycle — questions, orders, refunds, cancellations, complaints, and tickets.",
    "User accounts with persistent history make the system ready for real-world deployment.",
    "The analytics dashboard provides genuine operational insight for business decision-making.",
]

bullet_block(slide, ["- " + s for s in summary],
             0.55, 1.8, 12.2, 3.5,
             size=15, color=LIGHT, gap=Pt(20), dark=True)

text(slide, "What could be added next:",
     0.55, 5.6, 4.5, 0.45, size=13, bold=True, color=MID_GREY)
text(slide, "Urdu language support   |   Email and SMS notifications   |   Admin ticket management panel",
     0.55, 6.05, 12.0, 0.45, size=12, color=MID_GREY, italic=True)

text(slide, "Muhammad Huzaifa  |  2025",
     0.55, 6.7, 5.0, 0.4, size=11, color=MID_GREY)


# =============================================================================
# SAVE
# =============================================================================
out = "E-Commerce_AI_Support_Presentation.pptx"
prs.save(out)
print(f"Saved: {out}  ({len(prs.slides)} slides)")
